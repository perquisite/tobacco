# -*- coding: utf-8 -*-
import sys
import os
main_path = os.path.abspath(__file__)
root_path = os.path.dirname(os.path.dirname(main_path))
sys.path.append(root_path)
"""
Module implementing MainWindow.
"""
# 为防止pyinstaller打包exe报错，此处添加tobacco的路径
import re
import shutil

import cgitb
import docx
import xlwt
import xlrd
from PyQt5 import QtWidgets


from yancaoRegularDemo.Resource.Multi_Table.MultiTableProcessor import MultiTableProcessor
from yancaoRegularDemo.Resource.tools.get_pictures import get_pictures_single
from yancaoRegularDemo.Resource.tools.tangyuhao_function import get_desktop
from yancaoRegularDemo.lack_file import lack_file_dir

# sys.path.append("E:/Summer_Project/tobacco")

from PyQt5.QtCore import pyqtSlot, QBasicTimer
from PyQt5.QtWidgets import QMainWindow, QApplication, QMessageBox, QListView, QAbstractItemView, QTreeView
from PyQt5.QtWidgets import QFileDialog
# from past.builtins import raw_input

from yancaoRegularDemo.Resource.UiLayout import Ui_MainWindow

# from yancaoRegularDemo.Resource.xiejunyu import Precessor2
from yancaoRegularDemo.Resource.tangyuhao import Precessor1


class MessageBox(QtWidgets.QWidget):  # 继承自父类QtWidgets.QWidget
    def __init__(self, t, parent=None):  # parent = None代表此QWidget属于最上层的窗口,也就是MainWindows.
        QtWidgets.QWidget.__init__(self)  # 因为继承关系，要对父类初始化
        self.show_message(t)  # 信号槽

    def show_message(self, t):
        if t == 0:
            QtWidgets.QMessageBox.critical(self, "错误", "未在桌面检测到输入文件夹或输出文件夹")

        if t == 1:
            QtWidgets.QMessageBox.critical(self, "错误", "输入文件夹只能包含文件或只能包含文件夹")
class MainWindow(QMainWindow, Ui_MainWindow):
    """
    Class documentation goes here.
    """
    docs_to_process = None
    need_to_export = True
    export_dir = None
    status = "ready"  # 当前程序是否是就绪状态（无需要进行的任务）

    def __init__(self, parent=None):
        """
        Constructor

        @param parent reference to the parent widget (defaults to None)
        @type QWidget (optional)
        """
        super(MainWindow, self).__init__(parent)
        self.timer = QBasicTimer()
        self.setupUi(self)
        self.step = 0
        self.step_size = None
        self.flag = 0
        self.line = 0
        self.column = 0
        self.contract_check_result = []
        self.style_common = xlwt.XFStyle()
        self.export_dir = get_desktop() + '\\' + '烟草运行结果'
        self.lineEdit.setText(self.export_dir)
        self.compare.setVisible(False)

        my_dir = get_desktop() + '\\' + '烟草输入_文件夹'
        self.lineEdit_2.setText("输入路径：" + my_dir)
        if os.path.exists(self.export_dir) == False or os.path.exists(my_dir) == False:
            app = QtWidgets.QApplication(sys.argv)
            window = MessageBox(t=0)
            window.show()
            sys.exit(app.exec_())



    def timerEvent(self, a0: 'QTimerEvent') -> None:
        global doc1, file_p, worksheet, t, Multi, time0

        font = xlwt.Font()
        font.bold = False
        font.height = 12 * 20
        self.style_common.font = font

        if self.step >= self.step_size:
            print("全部审查完毕。。。")
            self.status = "ready"
            # self.step_size = -1  # 目的是避免重复点击“审查文件”按钮
            self.pushButton_2.setText('开始审查文件')
            # self.textBrowser.append("全部审查完毕。。。")
            QMessageBox.information(self, u"提示信息", u"全部审查完毕！")
            self.timer.stop()
            self.step = 0
            return

        self.status = "work"
        file = self.docs_to_process[self.step]
        print("-----------------------------------\n正在审查文件：" + file)
        self.textBrowser.append("-----------------------------------\n正在审查文件：" + file)

        if self.have_dir == 1:
            self.contract_check_result=[]
            Multi = MultiTableProcessor(input_dir_dictionary=self.dir_files, output_dir=self.lineEdit.text(),
                                        need_to_export=self.checkBox.isChecked(), progressBar=self.progressBar,
                                        textBrowser=self.textBrowser)
            self.contract_check_result = Multi.check()
            self.progressBar.setValue(int(100))
            self.step = self.step_size


        else:
            self.contract_check_result=[]
            d = Precessor1(file, self.export_dir, self.need_to_export)
            self.contract_check_result = self.contract_check_result + [file, d.action()]
            self.step = self.step + 1
            self.progressBar.setValue(int(self.step / self.step_size * 100))

        if self.need_to_export:
            # 判断输出文件是否存在，若不存在则创建文件文件，程序运行一次只判断一次
            import time
            time0 = time.strftime("%Y-%m-%d_%H_%M_%S", time.localtime())
            file_p = os.path.join(self.export_dir + '/' + "案卷审查结果表格_"+time0+".xls")
            if self.flag == 0:
                file_all = []
                for root, dirs, files in os.walk(self.export_dir):
                    for f in files:
                        file_all.append(os.path.join(self.export_dir + '/', f))
                if file_p not in file_all:
                    doc0 = xlwt.Workbook(encoding='utf-8')
                    worksheet = doc0.add_sheet('My Worksheet')
                    doc0.save(file_p)
                else:
                    print("已存在结果文件")
                self.flag = 1

        if self.need_to_export:
            # if self.chooseDir.isChecked():
            if self.have_dir == 1:
                doc1 = xlwt.Workbook(file_p)

                get_all_inf = Multi.get_all_info()
                worksheet_all = doc1.add_sheet("总览")
                first_col = worksheet_all.col(0)  # xlwt中是行和列都是从0开始计算的
                sec_col = worksheet_all.col(1)
                third_col = worksheet_all.col(2)
                forth_col = worksheet_all.col(3)

                first_col.width = 256 * 40
                sec_col.width = 256 * 20
                third_col.width = 256 * 20
                forth_col.width = 256 * 40

                self.line = 0

                style = xlwt.XFStyle()
                font = xlwt.Font()
                font.bold = True
                font.height = 18 * 20
                style.font = font

                worksheet_all.write(0, 0, "案卷名称", style)
                worksheet_all.write(0, 1, "实际审查项数量", style)
                worksheet_all.write(0, 2, "错误项数量", style)
                worksheet_all.write(0, 3, "需要人工审查的数量", style)
                self.line = self.line + 1
                for item in get_all_inf:
                    worksheet_all.write(self.line, 0, item[0], self.style_common)
                    worksheet_all.write(self.line, 1, item[1], self.style_common)
                    worksheet_all.write(self.line, 2, item[2], self.style_common)
                    worksheet_all.write(self.line, 3, item[3], self.style_common)
                    self.line = self.line + 1

                dir_name0 = None
                for item in self.contract_check_result:
                    if isinstance(item, str):
                        if "self." not in item:
                            pass
                        else:
                            self.contract_check_result.remove(item)
                            continue
                    elif isinstance(item, list):
                        item = [i for i in item if i[:5] != 'self.']
                    if "-----------------------------------\n正在审查文件：" in item:
                        pattern = r"-----------------------------------\n正在审查文件：(.*)"
                        t = re.findall(pattern, item)[0]  # t代表文件名

                        end_pos = t.rfind('/') - 1
                        dir_name = t[t.rfind('/', 1, end_pos) + 1:t.rfind('/', 1)]  # dir_name代表文件所在的文件夹
                        if dir_name0 == None:
                            self.line = 0
                            dir_name0 = dir_name  # 储存上一个文件夹
                            worksheet = doc1.add_sheet(dir_name)
                            first_col = worksheet.col(0)  # xlwt中是行和列都是从0开始计算的
                            sec_col = worksheet.col(1)

                            first_col.width = 256 * 70
                            sec_col.width = 256 * 100

                            style = xlwt.XFStyle()
                            font = xlwt.Font()
                            font.bold = True
                            font.height = 18 * 20
                            style.font = font

                            worksheet.write(0, 0, '文书路径', style)
                            worksheet.write(0, 1, '错误内容', style)
                            self.line = self.line + 1

                        if dir_name != dir_name0:  # 换sheet
                            worksheet = doc1.add_sheet(dir_name)
                            first_col = worksheet.col(0)  # xlwt中是行和列都是从0开始计算的
                            sec_col = worksheet.col(1)

                            first_col.width = 256 * 70
                            sec_col.width = 256 * 100

                            dir_name0 = dir_name
                            self.line = 0

                            style = xlwt.XFStyle()
                            font = xlwt.Font()
                            font.bold = True
                            font.height = 18 * 20
                            style.font = font

                            worksheet.write(0, 0, '文书路径', style)
                            worksheet.write(0, 1, '错误内容', style)
                            self.line = self.line + 1


                    elif "正在审查同案由案件比较信息：" in item:
                        self.line = 0
                        self.worksheet_compare = doc1.add_sheet("比较结果")
                        first_col = self.worksheet_compare.col(0)  # xlwt中是行和列都是从0开始计算的
                        sec_col = self.worksheet_compare.col(1)

                        first_col.width = 256 * 70
                        sec_col.width = 256 * 200

                        style = xlwt.XFStyle()
                        font = xlwt.Font()
                        font.bold = True
                        font.height = 18 * 20
                        style.font = font

                        self.worksheet_compare.write(0, 0, "比较项名称", style)
                        self.worksheet_compare.write(0, 1, "比较结果", style)
                        self.line = self.line + 1


                    elif "案件处理审批表_.docx\n" in item:
                        self.worksheet_compare.write(self.line, 0, item, self.style_common)

                    elif "同案由案件" in item:
                        self.worksheet_compare.write(self.line, 1, item, self.style_common)
                        self.line = self.line + 1

                    else:
                        for i in item:
                            if "×" in i:
                                i = i[2:]
                            worksheet.write(self.line, 0, t, self.style_common)
                            worksheet.write(self.line, 1, i, self.style_common)
                            self.line = self.line + 1

                doc1.save(file_p)
            else:
                file_p = os.path.join(self.export_dir + '/' + "案卷审查结果表格_"+time0+".xls")
                doc1 = xlwt.Workbook(file_p)
                for item in self.contract_check_result:
                    if isinstance(item, str):
                        if "self." not in item:
                            pass
                        else:
                            self.contract_check_result.remove(item)
                    elif isinstance(item, list):
                        item = [i for i in item if i[:5] != 'self.']
                    if ".docx" in item:
                        self.line = 0
                        worksheet = doc1.add_sheet(
                            item[item.rfind('/', 1) + 1:item.rfind('.', 1)].split("Desktop")[1].replace("\\", "_"))
                        first_col = worksheet.col(0)  # xlwt中是行和列都是从0开始计算的
                        sec_col = worksheet.col(1)

                        first_col.width = 256 * 70
                        sec_col.width = 256 * 100
                    else:
                        for i in item:
                            worksheet.write(self.line, 0, i, self.style_common)
                            self.line = self.line + 1
                doc1.save(file_p)

        # 输出提示信息
        # if self.chooseDir.isChecked() == True:
        if self.have_dir == 1:
            for item in self.contract_check_result:
                if "-----------------------------------\n正在审查文件：" in item:
                    self.textBrowser.append(item)
                elif "正在审查同案由案件比较信息：" in item:
                    self.textBrowser.append(item)
                elif "同案由案件" in item:
                    self.textBrowser.append(item)
                elif "案件处理审批表_.docx\n" in item:
                    self.textBrowser.append(item)
                else:
                    for i in item:
                        self.textBrowser.append(i)

        else:
            for item in self.contract_check_result:
                if ".docx" in item:
                    self.textBrowser.append(item)
                else:
                    for i in item:
                        self.textBrowser.append(i)

    # 识别文件 按钮
    @pyqtSlot()
    def on_pushButton_clicked(self):
        """
        Slot documentation goes here.
        """
        my_dir = get_desktop() + '\\' + '烟草输入_文件夹'
        self.dir_files = {}
        self.have_dir = 0
        l_all = os.listdir(my_dir)
        for l in l_all:
            if l != 'picture' and os.path.isdir(os.path.join(my_dir, l)) == True:
                self.have_dir = 1

        self.have_doc = 0
        for l in l_all:
            if l != 'picture' and os.path.isfile(os.path.join(my_dir, l)) == True:
                self.have_doc = 1

        if self.have_dir and self.have_doc:
            app = QtWidgets.QApplication(sys.argv)
            window = MessageBox(t=1)
            window.show()
            sys.exit(app.exec_())
        self.progressBar.setValue(0)


        if self.have_dir == 0:
            self.lineEdit_2.setText("输入路径：" + my_dir + "    以无文件夹模式启动")
            self.compare.setVisible(False)
            my_dir = get_desktop() + '\\' + '烟草输入_文件夹'
            # my_dir_path = QFileDialog.getExistingDirectory(self, u"打开文件夹", '/')
            my_file_path = []
            for root0, dirs0, files0 in os.walk(my_dir):
                for f in files0:
                    if "~$" not in f:
                        my_file_path.append(os.path.join(root0, f))
            print(my_file_path)
            if len(my_file_path) == 0:
                self.textBrowser_2.setText("您取消了操作！")
                return
            self.textBrowser_2.setText("您选择了以下文件：")
            for dir in my_file_path:
                self.textBrowser_2.append(dir)
            self.docs_to_process = my_file_path
            self.step_size = len(my_file_path)
        else:
            self.lineEdit_2.setText("输入路径：" + my_dir + "    以有文件夹模式启动")
            self.compare.setVisible(True)
            my_dir = get_desktop() + '\\' + '烟草输入_文件夹'
            # my_dir_path = QFileDialog.getExistingDirectory(self, u"打开文件夹", '/')
            my_file_path = []
            for root0, dirs0, files0 in os.walk(my_dir):
                for dir in dirs0:
                    x = os.path.join(root0, dir)
                    for root1, dirs1, files1 in os.walk(x):
                        for f in files1:
                            path = os.path.join(root1, f)
                            if "jpeg" in path or "png" in path or "jpg" in path or "~" in path:
                                continue
                            else:
                                my_file_path.append(os.path.join(root1, f))

            dir_now = my_file_path[0].split("烟草输入_文件夹\\")[1].split("\\")[0]
            my_file_path_single = []
            for f in my_file_path:
                dir = f.split("烟草输入_文件夹\\")[1].split("\\")[0]
                f = f.replace("\\", "/")
                if dir_now == dir:
                    my_file_path_single.append(f)
                else:
                    self.dir_files[(get_desktop() + '\\' + dir_now).replace("\\", "/")] = my_file_path_single
                    dir_now = dir
                    my_file_path_single = []
                    my_file_path_single.append(f)
            self.dir_files[(get_desktop() + '\\' + dir_now).replace("\\", "/")] = my_file_path_single
            # print(self.dir_files)
            folders = []
            for root, dirs, files in os.walk(my_dir):
                if 'picture' not in root:
                    folders.append(os.path.join(root, ))
            if len(folders) == 0:
                self.textBrowser_2.setText("您取消了操作！")
                return

            real_folders = []
            # 防止父文件夹被选择
            for i in range(0, len(folders) - 1):
                if folders[i] not in folders[i + 1]:
                    real_folders.append(folders[i])
            real_folders.append(folders[-1])
            # print(real_folders)

            if len(my_file_path) == 0:
                self.textBrowser_2.setText("文件夹可使用文件为空！")
                return

            self.textBrowser_2.setText("您选择了以下文件夹以及文件：")
            self.textBrowser_2.append('------------文件夹------------')

            for dir in real_folders:
                self.textBrowser_2.append(dir)
                get_pictures_single(dir)
            self.textBrowser_3.setText("以下文件夹缺失文件：")
            for dir in real_folders:
                list = lack_file_dir(dir)
                self.textBrowser_3.append('------------文件夹------------')
                self.textBrowser_3.append(dir)
                self.textBrowser_3.append('-----------缺失文件-----------')
                for l in list:
                    self.textBrowser_3.append(l)
                self.textBrowser_3.append("\n")
                self.textBrowser_3.append("\n")
            self.textBrowser_2.append('-------------文件-------------')
            for dir in my_file_path:
                self.textBrowser_2.append(dir)
            self.docs_to_process = my_file_path
            self.step_size = len(my_file_path)

    # "是否进行比较"一系列 按钮按下
    @pyqtSlot()
    def on_compare_clicked(self):
        """
        Slot documentation goes here.
        """
        if self.compare.isChecked() == True:

            self.compare.setText("进行比较")
        else:
            self.compare.setText("是否进行比较")

    @pyqtSlot()
    def on_pushButton_2_clicked(self):
        """
        Slot documentation goes here.
        """
        if self.status == "ready":
            self.textBrowser.clear()

            if self.docs_to_process == None:
                QMessageBox.information(self, u"提示信息", u"请添加要审查的文件！")
                return
            # if self.step_size == -1:
            #     # print("任务已经完成！无需重复点击审查。\n若需要审查新的文件,请重新选择待审查的新文件")
            #     # self.textBrowser.append("任务已经完成！无需重复点击审查。\n若需要审查新的文件,请重新选择待审查的新文件")
            #     QMessageBox.information(self, u"提示信息", u"任务已经完成！无需重复点击审查。\n若需要审查新的文件,请重新选择待审查的新文件")
            #     return

        if self.timer.isActive():
            self.textBrowser.append("#####################已暂停！")
            self.timer.stop()
            self.pushButton_2.setText('继续')
        else:
            self.textBrowser.append("已开始！")
            self.timer.start(100, self)
            self.pushButton_2.setText('暂停')

    # "导出审查结果文件" CheckBox按下
    @pyqtSlot()
    def on_checkBox_clicked(self):
        """
        Slot documentation goes here.
        """
        self.need_to_export = self.checkBox.isChecked()
        if self.checkBox.isChecked():
            self.lineEdit.setVisible(True)
            self.pushButton_4.setVisible(True)
        else:
            self.lineEdit.setVisible(False)
            self.pushButton_4.setVisible(False)

    @pyqtSlot()
    def on_pushButton_3_clicked(self):
        """
        Slot documentation goes here.
        """
        my_dir = get_desktop() + '\\' + '烟草输入_文件夹'
        if os.path.exists(my_dir):
            os.system("explorer.exe %s" % my_dir)
        else:
            return

    @pyqtSlot()
    def on_pushButton_4_clicked(self):
        """
        Slot documentation goes here.
        """
        my_dir = get_desktop() + '\\' + '烟草运行结果'
        if os.path.exists(my_dir):
            os.system("explorer.exe %s" % my_dir)
        else:
            return


if __name__ == "__main__":
    import sys

    cgitb.enable(format='text')
    app = QApplication(sys.argv)
    ui = MainWindow()
    ui.show()
    sys.exit(app.exec_())
